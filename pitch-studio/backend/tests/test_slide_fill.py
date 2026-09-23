"""Stage 4 tests: fill, gate, cache, render, persist and degrade-to-brand.

Everything external is mocked — the LLM slot-fill, the Opus vision review, the
slide renderer, brand-page fetch and object storage — so the whole pipeline is
exercised deterministically without a browser or network. The database is a real
in-memory SQLite so the shared cache, uniqueness and race handling are genuine.
"""

from __future__ import annotations

import io
import json
import unittest
from types import SimpleNamespace
from unittest import mock

from sqlalchemy import create_engine
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.database import Base
from backend.generated_slides import (
    compute_claim_hash,
    compute_render_hash,
    generated_slide_key,
)
from backend.models import GeneratedSlide, GeneratedSlideAttempt, Generation, User
from backend.pipeline.brand_deck import BrandSlide, render_pptx
from backend.pipeline.deck import GENERATED_SOURCE, deck_slide_from_planned
from backend.pipeline.gaps import (
    CLOSING_PAGE,
    COVER_PAGE,
    GeneratedSlidePlaceholder,
    build_gap_plan,
)
from backend.pipeline.slide_fill import (
    ApprovedPhoto,
    GeneratedSlideInstance,
    _fill_payload,
    _prepare_photos,
    _repair_budgets,
    _resolve_values,
    _run_vision,
    apply_deck_conventions,
    gate_budgets,
    gate_copy_register,
    gate_number_provenance,
    gate_schema,
    make_recommended_photo_fn,
    realize_generated_slides,
)
from backend.pipeline.slide_templates import template_spec

SLIDE_MODULE = "backend.pipeline.slide_fill"


def make_jpeg(color=(200, 60, 40)) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (64, 36), color).save(buffer, format="JPEG")
    return buffer.getvalue()


def fact(fid: int, name: str, value: str, status: str = "verified") -> SimpleNamespace:
    return SimpleNamespace(id=fid, fact=name, value=value, status=status)


def passage(asset_id: int, text: str, module_ids: str = "") -> dict:
    return {
        "asset_id": asset_id,
        "title": "Report",
        "module_ids": module_ids,
        "start_page": 1,
        "end_page": 1,
        "text": text,
    }


class FakeRenderer:
    """A stand-in for the warm slide renderer with scripted outcomes."""

    def __init__(self, jpeg: bytes | None = None, raises=None):
        self.jpeg = jpeg if jpeg is not None else make_jpeg()
        self.raises = raises
        self.calls: list[dict] = []
        self.photo_calls: list[dict] = []

    def render(self, manifest, values, *, photos=None):
        self.calls.append(dict(values))
        self.photo_calls.append(dict(photos or {}))
        if self.raises is not None:
            raise self.raises
        return SimpleNamespace(jpeg=self.jpeg, template_id=manifest.template_id, fitted_sizes={})


def placeholder(**overrides) -> GeneratedSlidePlaceholder:
    base = dict(
        slide_key="generated:m13-what-we-are-not",
        gap_kind="context",
        claim="We are not a glorified weekend bootcamp; we are a real institution",
        title="What we are not",
        module_id="M13",
        template_id="section-divider-light",
        tone="light",
        subtitle="",
        recipe_modules=("M13",),
        source_fact_ids=(),
        replaced_page=51,
        replaced_module_id="M04",
        replaced_label="The five Outclass challenges",
    )
    base.update(overrides)
    return GeneratedSlidePlaceholder(**base)


class FillPayloadTests(unittest.TestCase):
    def test_script_excerpt_is_the_content_source(self):
        payload = _fill_payload(
            template_spec("section-divider-light"),
            placeholder(),
            "The approved script explains the institution, not a weekend bootcamp.",
            [],
            [],
            [],
        )

        self.assertIn("approved script explains", payload["approved_script_excerpt"])
        self.assertIn(
            "Infer the slide's subject and wording only from approved_script_excerpt.",
            payload["conventions"],
        )
        self.assertIn(
            "Use style_exemplars only for tone, length and structure; never copy their subject matter.",
            payload["conventions"],
        )


class MemoryDbTestCase(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine)
        self.db = self.Session()

    def tearDown(self):
        self.db.close()


# ---------------------------------------------------------------------------
# Deterministic hashes
# ---------------------------------------------------------------------------


class ClaimHashTests(unittest.TestCase):
    def _hash(self, *, value="₹27.78 LPA median", status="verified", passages=()):
        return compute_claim_hash(
            claim="Placements are strong",
            module_id="M07",
            template_id="section-divider-light",
            tone="light",
            source_facts=[fact(1, "Median CTC", value, status)],
            source_passages=passages,
        )

    def test_is_stable_for_identical_inputs(self):
        self.assertEqual(self._hash(), self._hash())

    def test_changes_when_locked_fact_content_changes(self):
        self.assertNotEqual(self._hash(value="₹27.78 LPA median"), self._hash(value="₹30.00 LPA median"))

    def test_changes_when_fact_verification_status_changes(self):
        self.assertNotEqual(self._hash(status="verified"), self._hash(status="needs_source"))

    def test_changes_when_a_source_passage_changes(self):
        one = self._hash(passages=[passage(9, "200 startups, ₹25 Cr raised", "M07")])
        two = self._hash(passages=[passage(9, "210 startups, ₹25 Cr raised", "M07")])
        self.assertNotEqual(one, two)

    def test_is_order_independent_across_facts(self):
        facts_a = [fact(1, "A", "1"), fact(2, "B", "2")]
        facts_b = [fact(2, "B", "2"), fact(1, "A", "1")]
        common = dict(claim="c", module_id="M01", template_id="section-divider-light", tone="light")
        self.assertEqual(
            compute_claim_hash(**common, source_facts=facts_a),
            compute_claim_hash(**common, source_facts=facts_b),
        )


class RenderHashTests(unittest.TestCase):
    def _hash(self, *, slots=None, version="1", font="fonts-1"):
        return compute_render_hash(
            template_id="section-divider-light",
            template_version=version,
            slot_values=slots or {"title": "Not a *bootcamp*", "subtitle": ""},
            font_bundle_version=font,
        )

    def test_stable_for_identical_inputs(self):
        self.assertEqual(self._hash(), self._hash())

    def test_changes_with_slot_copy(self):
        self.assertNotEqual(self._hash(), self._hash(slots={"title": "A different line", "subtitle": ""}))

    def test_changes_with_template_version(self):
        self.assertNotEqual(self._hash(version="1"), self._hash(version="2"))

    def test_changes_with_font_bundle_version(self):
        self.assertNotEqual(self._hash(font="fonts-1"), self._hash(font="fonts-2"))

    def test_slide_key_is_derived_and_urlsafe(self):
        digest = self._hash()
        key = generated_slide_key(digest)
        self.assertTrue(key.startswith("generated-"))
        self.assertNotIn(":", key)

    def _photo_hash(self, *, asset_id=101, focus_x=0.5, rect=None):
        crop = {
            "fit": "cover",
            "focus_x": focus_x,
            "focus_y": 0.5,
            "rect": rect or [0, 0, 1920, 1080],
            "rotation_deg": 0.0,
        }
        return compute_render_hash(
            template_id="campus-photo-dark",
            template_version="1",
            slot_values={"headline": "Classrooms", "caption": ""},
            font_bundle_version="fonts-1",
            photo={"photos": [{"slot": "hero_photo", "asset_id": asset_id, "key": "k", "crop": crop}]},
        )

    def test_photo_presence_changes_the_render_hash(self):
        no_photo = compute_render_hash(
            template_id="campus-photo-dark", template_version="1",
            slot_values={"headline": "Classrooms", "caption": ""}, font_bundle_version="fonts-1",
        )
        self.assertNotEqual(no_photo, self._photo_hash())

    def test_render_hash_changes_when_photo_asset_changes(self):
        self.assertNotEqual(self._photo_hash(asset_id=101), self._photo_hash(asset_id=202))

    def test_render_hash_changes_when_crop_changes(self):
        self.assertNotEqual(self._photo_hash(focus_x=0.5), self._photo_hash(focus_x=0.25))
        self.assertNotEqual(
            self._photo_hash(rect=[0, 0, 1920, 1080]),
            self._photo_hash(rect=[100, 0, 800, 600]),
        )


# ---------------------------------------------------------------------------
# Gates and conventions (pure)
# ---------------------------------------------------------------------------


class ConventionAndGateTests(unittest.TestCase):
    def setUp(self):
        self.spec = template_spec("section-divider-light")

    def test_conventions_apply_ampersand_and_strip_terminal_period(self):
        self.assertEqual(apply_deck_conventions("Learn and build."), "Learn & build")
        self.assertEqual(apply_deck_conventions("Immersions"), "Immersions")
        self.assertEqual(apply_deck_conventions("What next..."), "What next...")

    def test_schema_gate_requires_title(self):
        self.assertTrue(gate_schema(self.spec, {"title": "  ", "subtitle": "x"}))
        self.assertFalse(gate_schema(self.spec, {"title": "Real title", "subtitle": ""}))

    def test_budget_gate_flags_over_length_copy(self):
        long_title = "word " * 40
        violations = gate_budgets(self.spec, {"title": long_title, "subtitle": ""})
        self.assertTrue(any("characters" in v or "words" in v for v in violations))

    def test_budget_gate_passes_short_copy(self):
        self.assertFalse(gate_budgets(self.spec, {"title": "Not a *bootcamp*", "subtitle": ""}))

    def test_budget_gate_ignores_emphasis_markup_in_character_count(self):
        spec = template_spec("stat-light")
        visible = "x" * 48
        self.assertFalse(gate_budgets(spec, {"title": f"*{visible}*", "hero_value": "3x", "hero_label": "Salary"}))

    def test_number_gate_rejects_unsupported_numeral(self):
        allowed = {"27.78"}
        violations = gate_number_provenance(
            self.spec, {"title": "We place 500 students", "subtitle": ""}, allowed
        )
        self.assertTrue(any("500" in v for v in violations))

    def test_number_gate_accepts_supported_numeral(self):
        allowed = {"27.78", "500"}
        self.assertFalse(
            gate_number_provenance(
                self.spec, {"title": "Median is 27.78 LPA", "subtitle": ""}, allowed
            )
        )

    def test_register_gate_rejects_conservative_officialese(self):
        phrases = [
            "following Government of Haryana approval",
            "As per the approved policy",
            "Pursuant to the new framework",
            "In order to build practical judgment",
            "With respect to student outcomes",
            "For the purpose of industry readiness",
            "In this regard, learning stays practical",
            "It should be noted that founders teach here",
            "Admissions are hereby opened",
        ]
        for phrase in phrases:
            with self.subTest(phrase=phrase):
                self.assertTrue(
                    gate_copy_register(self.spec, {"title": phrase, "subtitle": ""})
                )

    def test_divider_register_requires_an_italic_serif_accent(self):
        violations = gate_copy_register(
            self.spec,
            {
                "title": "What happens to ventures that fail?",
                "subtitle": "",
            },
        )
        self.assertTrue(
            any("italic-serif accent" in violation for violation in violations)
        )
        self.assertFalse(
            gate_copy_register(
                self.spec,
                {
                    "title": "What happens to *ventures that fail?*",
                    "subtitle": "",
                },
            )
        )

    def test_divider_question_drops_subtitle_to_avoid_two_line_collision(self):
        values = _repair_budgets(
            self.spec,
            _resolve_values(
                self.spec,
                {
                    "title": "What happens to students whose *ventures fail*?",
                    "subtitle": "Failure teaches reflection",
                },
            ),
            {},
            final_attempt=False,
        )
        self.assertEqual(values["subtitle"], "")

    def test_short_divider_title_keeps_subtitle(self):
        values = _repair_budgets(
            self.spec,
            _resolve_values(
                self.spec,
                {
                    "title": "*Immersions*",
                    "subtitle": "How students *learn by traveling*",
                },
            ),
            {},
            final_attempt=False,
        )
        self.assertEqual(values["subtitle"], "How students *learn by traveling*")

    def test_register_gate_accepts_all_real_brand_exemplars(self):
        for template_id in (
            "section-divider-light",
            "stat-light",
            "campus-photo-dark",
            "programme-list-light",
        ):
            spec = template_spec(template_id)
            for exemplar in spec.exemplars:
                with self.subTest(template_id=template_id, exemplar=exemplar):
                    values = {name: "" for name in spec.fillable_slots}
                    values.update({name: [] for name in spec.fillable_lists})
                    if spec.fillable_slots:
                        values[spec.fillable_slots[0]] = exemplar
                    else:
                        values[spec.fillable_lists[0]] = [exemplar]
                    self.assertFalse(gate_copy_register(spec, values))


class VisionContractTests(unittest.TestCase):
    def test_structured_copy_and_furniture_violations_are_preserved(self):
        review = _run_vision(
            lambda *_: {
                "passed": False,
                "violations": [
                    {"category": "furniture", "detail": "Missing lockup"},
                    {"category": "copy", "detail": "Awkward caption"},
                ],
            },
            b"jpeg",
            [],
            template_spec("section-divider-light"),
            placeholder(),
        )
        self.assertEqual(
            review["violations"],
            [
                {"category": "furniture", "detail": "Missing lockup"},
                {"category": "copy", "detail": "Awkward caption"},
            ],
        )

    def test_legacy_string_violations_default_to_copy(self):
        review = _run_vision(
            lambda *_: {"passed": False, "violations": ["Awkward caption"]},
            b"jpeg",
            [],
            template_spec("section-divider-light"),
            placeholder(),
        )
        self.assertEqual(
            review["violations"],
            [{"category": "copy", "detail": "Awkward caption"}],
        )


# ---------------------------------------------------------------------------
# End-to-end realise (cache / miss / gates / degrade / persist)
# ---------------------------------------------------------------------------


class RealizeTests(MemoryDbTestCase):
    def _realize(self, plan, *, fill_fn, vision_fn, renderer, facts=None, passages=None, save_key="generated-slides/x.jpg", file_exists=True):
        with mock.patch(f"{SLIDE_MODULE}.save_generated_slide_image", return_value=save_key) as save, mock.patch(
            f"{SLIDE_MODULE}.file_exists", return_value=file_exists
        ):
            result = realize_generated_slides(
                self.db,
                plan,
                facts=facts if facts is not None else [],
                passages=passages if passages is not None else [],
                fill_fn=fill_fn,
                vision_fn=vision_fn,
                renderer=renderer,
                brand_page_fn=lambda page: make_jpeg((10, 10, 10)),
            )
        return result, save

    def test_no_placeholders_is_an_untouched_passthrough(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), BrandSlide(6, "M01", "Origin"), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        # A renderer/fill that would explode if ever touched.
        boom = mock.Mock(side_effect=AssertionError("must not run"))
        result = realize_generated_slides(
            self.db, plan, facts=[], passages=[], fill_fn=boom, vision_fn=boom, renderer=mock.Mock(render=boom)
        )
        self.assertEqual(result, plan)
        boom.assert_not_called()

    def test_miss_then_success_persists_and_returns_instance(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        renderer = FakeRenderer()
        fill = mock.Mock(return_value={"title": "Not a *weekend bootcamp*", "subtitle": ""})
        vision = mock.Mock(return_value={"passed": True, "violations": []})

        result, save = self._realize(plan, fill_fn=fill, vision_fn=vision, renderer=renderer)

        slide = result[1]
        self.assertIsInstance(slide, GeneratedSlideInstance)
        self.assertEqual(slide.source, GENERATED_SOURCE)
        self.assertIsNone(slide.page)
        self.assertTrue(slide.image_url.startswith("/api/generated-slides/generated-"))
        self.assertTrue(slide.image_url.endswith(".jpg"))
        self.assertEqual(slide.title, "Not a weekend bootcamp")  # emphasis markup stripped
        save.assert_called_once()
        fill.assert_called_once()
        vision.assert_called_once()

        rows = self.db.query(GeneratedSlide).all()
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row.template_id, "section-divider-light")
        self.assertEqual(row.tone, "light")
        self.assertEqual(row.file_key, "generated-slides/x.jpg")
        self.assertEqual(row.status, "ready")
        self.assertEqual(json.loads(row.slot_values_json)["title"], "Not a *weekend bootcamp*")
        self.assertEqual(row.slide_key, generated_slide_key(row.render_hash))

    def test_budget_retry_keeps_slots_that_already_fitted(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        long_subtitle = "Hands-on learning with founders " * 4
        long_title = "Not a bootcamp " * 10
        fill = mock.Mock(side_effect=[
            {"title": "Not a *bootcamp*", "subtitle": long_subtitle},
            {"title": long_title, "subtitle": "Built by founders"},
        ])
        vision = mock.Mock(return_value={"passed": True, "violations": []})

        result, save = self._realize(plan, fill_fn=fill, vision_fn=vision, renderer=FakeRenderer())

        self.assertIsInstance(result[1], GeneratedSlideInstance)
        self.assertEqual(fill.call_count, 2)
        retry_payload = fill.call_args_list[1].args[0]
        self.assertEqual(retry_payload["previous_values"]["title"], "Not a *bootcamp*")
        row = self.db.query(GeneratedSlide).one()
        self.assertEqual(
            json.loads(row.slot_values_json),
            {"title": "Not a *bootcamp*", "subtitle": "Built by founders"},
        )

    def test_droppable_line_still_over_budget_is_blanked_on_final_attempt(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        long_subtitle = "Hands-on learning with founders " * 4
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": long_subtitle})
        vision = mock.Mock(return_value={"passed": True, "violations": []})

        result, save = self._realize(plan, fill_fn=fill, vision_fn=vision, renderer=FakeRenderer())

        self.assertIsInstance(result[1], GeneratedSlideInstance)
        self.assertEqual(fill.call_count, 3)
        row = self.db.query(GeneratedSlide).one()
        self.assertEqual(json.loads(row.slot_values_json)["subtitle"], "")

    def test_cache_hit_reuses_without_model_vision_or_render(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""})
        vision = mock.Mock(return_value={"passed": True, "violations": []})
        renderer = FakeRenderer()

        # First pass populates the shared cache.
        self._realize(plan, fill_fn=fill, vision_fn=vision, renderer=renderer)
        self.assertEqual(fill.call_count, 1)

        # Second pass (a different user, fresh plan) must hit the cache.
        fill2 = mock.Mock(side_effect=AssertionError("model must not run on a cache hit"))
        vision2 = mock.Mock(side_effect=AssertionError("vision must not run on a cache hit"))
        renderer2 = FakeRenderer(raises=AssertionError("render must not run on a cache hit"))
        result, save2 = self._realize(
            [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")],
            fill_fn=fill2,
            vision_fn=vision2,
            renderer=renderer2,
        )
        self.assertIsInstance(result[1], GeneratedSlideInstance)
        save2.assert_not_called()
        self.assertEqual(self.db.query(GeneratedSlide).count(), 1)

    def test_cache_row_with_missing_pixels_is_a_miss(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        # Seed a matching row, but its object is gone from storage.
        p = placeholder()
        claim_hash = compute_claim_hash(
            claim=p.claim, module_id=p.module_id, template_id=p.template_id, tone=p.tone,
        )
        self.db.add(
            GeneratedSlide(
                slide_key="generated-stale", claim_hash=claim_hash, render_hash="a" * 64,
                template_id=p.template_id, tone="light", file_key="generated-slides/gone.jpg", status="ready",
            )
        )
        self.db.commit()
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""})
        vision = mock.Mock(return_value={"passed": True, "violations": []})
        result, _ = self._realize(
            plan, fill_fn=fill, vision_fn=vision, renderer=FakeRenderer(), file_exists=False
        )
        # Regenerated rather than serving a dead URL.
        fill.assert_called_once()
        self.assertIsInstance(result[1], GeneratedSlideInstance)

    def test_unsupported_number_is_rejected_and_degrades(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(source_fact_ids=(1,)), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        renderer = FakeRenderer()
        # The model keeps inventing an unsupported number.
        fill = mock.Mock(return_value={"title": "We place 999 students", "subtitle": ""})
        vision = mock.Mock(return_value={"passed": True, "violations": []})
        facts = [fact(1, "Median CTC", "₹27.78 LPA")]

        result, save = self._realize(plan, fill_fn=fill, vision_fn=vision, renderer=renderer, facts=facts)

        restored = result[1]
        self.assertIsInstance(restored, BrandSlide)
        self.assertEqual(restored.page, 51)  # the displaced brand page, restored
        self.assertEqual(fill.call_count, 3)  # three retryable copy attempts
        renderer_calls = renderer.calls
        self.assertEqual(renderer_calls, [])  # number gate precedes the render
        vision.assert_not_called()
        save.assert_not_called()
        self.assertEqual(self.db.query(GeneratedSlide).count(), 0)

    def test_overflow_degrades_to_original_brand_slide(self):
        from backend.pipeline.slide_render import SlideOverflowError

        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        renderer = FakeRenderer(raises=SlideOverflowError(["title"]))
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""})
        vision = mock.Mock(return_value={"passed": True, "violations": []})

        result, save = self._realize(plan, fill_fn=fill, vision_fn=vision, renderer=renderer)

        self.assertIsInstance(result[1], BrandSlide)
        self.assertEqual(result[1].page, 51)
        self.assertEqual(fill.call_count, 3)
        vision.assert_not_called()  # overflow precedes vision
        save.assert_not_called()

    def test_vision_failure_then_pass_succeeds_on_retry(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        renderer = FakeRenderer()
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""})
        vision = mock.Mock(side_effect=[
            {"passed": False, "violations": ["Copy sits too high"]},
            {"passed": True, "violations": []},
        ])

        result, save = self._realize(plan, fill_fn=fill, vision_fn=vision, renderer=renderer)

        self.assertIsInstance(result[1], GeneratedSlideInstance)
        self.assertEqual(fill.call_count, 2)
        self.assertEqual(vision.call_count, 2)
        save.assert_called_once()

    def test_pure_copy_vision_failure_uses_all_three_attempts(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        renderer = FakeRenderer()
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""})
        vision = mock.Mock(return_value={
            "passed": False,
            "violations": [{"category": "copy", "detail": "Caption is bureaucratic"}],
        })

        result, save = self._realize(plan, fill_fn=fill, vision_fn=vision, renderer=renderer)

        self.assertIsInstance(result[1], BrandSlide)
        self.assertEqual(result[1].page, 51)
        self.assertEqual(fill.call_count, 3)
        self.assertEqual(vision.call_count, 3)
        self.assertEqual(
            fill.call_args_list[1].args[0]["corrections"],
            ["Caption is bureaucratic"],
        )
        save.assert_not_called()

    def test_furniture_vision_failure_degrades_immediately(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        renderer = FakeRenderer()
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""})
        vision = mock.Mock(return_value={
            "passed": False,
            "violations": [
                {"category": "furniture", "detail": "Masters' Union lockup is missing"},
                {"category": "copy", "detail": "Caption is bureaucratic"},
            ],
        })

        result, save = self._realize(plan, fill_fn=fill, vision_fn=vision, renderer=renderer)

        self.assertIsInstance(result[1], BrandSlide)
        self.assertEqual(fill.call_count, 1)
        self.assertEqual(vision.call_count, 1)
        self.assertEqual(len(renderer.calls), 1)
        save.assert_not_called()

    def test_vision_exception_is_treated_as_a_failed_gate(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        renderer = FakeRenderer()
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""})
        vision = mock.Mock(side_effect=RuntimeError("model down"))

        result, _ = self._realize(plan, fill_fn=fill, vision_fn=vision, renderer=renderer)

        self.assertIsInstance(result[1], BrandSlide)
        self.assertEqual(vision.call_count, 3)


def campus_placeholder(**overrides) -> GeneratedSlidePlaceholder:
    base = dict(
        template_id="campus-photo-dark",
        tone="dark",
        claim="A high-energy classroom moment on campus",
        title="Classrooms",
    )
    base.update(overrides)
    return placeholder(**base)


def approved_photo(asset_id=101, key="assets/campus.jpg") -> ApprovedPhoto:
    return ApprovedPhoto(asset_id=asset_id, key=key, data=make_jpeg((30, 120, 200)))


# A valid caption in the brand band (7-11 words / 43-72 chars).
_CAMPUS_CAPTION = "High-energy learning spaces with digital boards & AV integration"


class PhotoSourcingTests(unittest.TestCase):
    """`_prepare_photos` assigns approved photos and records provenance/crop."""

    def test_required_slot_is_filled_with_descriptor_and_asset_id(self):
        spec = template_spec("campus-photo-dark")
        photos, descriptors, asset_ids = _prepare_photos(
            spec, campus_placeholder(), lambda ph, sp: [approved_photo(asset_id=101)]
        )
        self.assertIn("hero_photo", photos)
        self.assertEqual(asset_ids, [101])
        self.assertEqual(descriptors[0]["asset_id"], 101)
        self.assertEqual(descriptors[0]["slot"], "hero_photo")
        self.assertEqual(descriptors[0]["key"], "assets/campus.jpg")
        self.assertIn("crop", descriptors[0])
        self.assertEqual(descriptors[0]["crop"]["fit"], "cover")

    def test_no_photo_available_leaves_slots_unfilled(self):
        spec = template_spec("campus-photo-dark")
        photos, descriptors, asset_ids = _prepare_photos(spec, campus_placeholder(), lambda ph, sp: [])
        self.assertEqual((photos, descriptors, asset_ids), ({}, [], []))

    def test_optional_photo_slots_take_photos_in_order(self):
        spec = template_spec("programme-list-light")  # two optional photo slots
        photos, descriptors, asset_ids = _prepare_photos(
            spec, placeholder(template_id="programme-list-light", tone="light"),
            lambda ph, sp: [approved_photo(asset_id=1)],  # only one available
        )
        self.assertEqual(list(photos), ["photo_1"])  # first slot only
        self.assertEqual(asset_ids, [1])

    def test_template_without_photo_slots_sources_nothing(self):
        spec = template_spec("section-divider-light")
        self.assertEqual(_prepare_photos(spec, placeholder(), lambda ph, sp: [approved_photo()]), ({}, [], []))


class RecommendedPhotoFnTests(unittest.TestCase):
    def test_blank_recipe_ref_yields_no_photos(self):
        fn = make_recommended_photo_fn(object(), "")
        self.assertEqual(fn(campus_placeholder(), template_spec("campus-photo-dark")), [])

    def test_template_without_photo_slots_yields_no_photos(self):
        fn = make_recommended_photo_fn(object(), "recipe-1")
        self.assertEqual(fn(placeholder(), template_spec("section-divider-light")), [])


class CampusRealizeTests(MemoryDbTestCase):
    def _realize(self, plan, *, fill_fn, vision_fn, renderer, photo_fn, save_key="generated-slides/c.jpg"):
        with mock.patch(f"{SLIDE_MODULE}.save_generated_slide_image", return_value=save_key) as save, mock.patch(
            f"{SLIDE_MODULE}.file_exists", return_value=True
        ):
            result = realize_generated_slides(
                self.db, plan, facts=[], passages=[],
                fill_fn=fill_fn, vision_fn=vision_fn, renderer=renderer,
                brand_page_fn=lambda page: make_jpeg((10, 10, 10)),
                photo_fn=photo_fn,
            )
        return result, save

    def test_photo_is_embedded_and_provenance_recorded(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), campus_placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        renderer = FakeRenderer()
        fill = mock.Mock(return_value={"headline": "Classrooms", "caption": _CAMPUS_CAPTION})
        vision = mock.Mock(return_value={"passed": True, "violations": []})

        result, save = self._realize(
            plan, fill_fn=fill, vision_fn=vision, renderer=renderer,
            photo_fn=lambda ph, sp: [approved_photo(asset_id=101)],
        )

        slide = result[1]
        self.assertIsInstance(slide, GeneratedSlideInstance)
        self.assertEqual(slide.title, "Classrooms")
        # The approved photo bytes reached the renderer's photo slot.
        self.assertIn("hero_photo", renderer.photo_calls[-1])
        save.assert_called_once()
        row = self.db.query(GeneratedSlide).one()
        self.assertIn("101", (row.source_asset_ids or "").split(","))

    def test_required_photo_missing_degrades_before_any_fill(self):
        plan = [BrandSlide(COVER_PAGE, "", "Cover"), campus_placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        renderer = FakeRenderer()
        fill = mock.Mock(return_value={"headline": "Classrooms", "caption": _CAMPUS_CAPTION})
        vision = mock.Mock(return_value={"passed": True, "violations": []})

        result, save = self._realize(
            plan, fill_fn=fill, vision_fn=vision, renderer=renderer,
            photo_fn=lambda ph, sp: [],  # no approved photo for the recipe
        )

        self.assertIsInstance(result[1], BrandSlide)
        self.assertEqual(result[1].page, 51)  # the displaced brand page, restored
        fill.assert_not_called()  # photo gate precedes fill/render entirely
        save.assert_not_called()
        self.assertEqual(self.db.query(GeneratedSlide).count(), 0)

    def test_persisted_render_hash_incorporates_the_photo(self):
        # The photo is part of the pixels: the render hash the campus slide is
        # stored under must differ from the hash the identical copy would get
        # with no photo. (The claim cache itself is keyed on the claim, which is
        # photo-independent — a photo is deterministic per recipe.)
        from backend.pipeline.slide_fill import FONT_BUNDLE_VERSION

        plan = [BrandSlide(COVER_PAGE, "", "Cover"), campus_placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        renderer = FakeRenderer()
        fill = mock.Mock(return_value={"headline": "Classrooms", "caption": _CAMPUS_CAPTION})
        vision = mock.Mock(return_value={"passed": True, "violations": []})

        self._realize(
            plan, fill_fn=fill, vision_fn=vision, renderer=renderer,
            photo_fn=lambda ph, sp: [approved_photo(asset_id=101)],
        )
        row = self.db.query(GeneratedSlide).one()
        no_photo_hash = compute_render_hash(
            template_id=row.template_id,
            template_version=row.template_version,
            slot_values=json.loads(row.slot_values_json),
            font_bundle_version=FONT_BUNDLE_VERSION,
            photo=None,
        )
        self.assertNotEqual(row.render_hash, no_photo_hash)


class PersistenceRaceTests(MemoryDbTestCase):
    def test_duplicate_render_hash_is_rejected_by_the_unique_constraint(self):
        common = dict(
            claim_hash="c" * 64, render_hash="d" * 64, template_id="section-divider-light",
            tone="light", file_key="k", status="ready",
        )
        self.db.add(GeneratedSlide(slide_key="generated-a", **common))
        self.db.commit()
        self.db.add(GeneratedSlide(slide_key="generated-b", **common))
        with self.assertRaises(IntegrityError):
            self.db.commit()
        self.db.rollback()

    def test_concurrent_realize_converges_on_one_shared_row(self):
        # Two "users" render the same claim; the second insert hits the unique
        # render_hash and reuses the row the first writer created.
        plan_a = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        plan_b = [BrandSlide(COVER_PAGE, "", "Cover"), placeholder(), BrandSlide(CLOSING_PAGE, "M14", "Close")]
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""})
        vision = mock.Mock(return_value={"passed": True, "violations": []})

        from backend.pipeline import slide_fill as sf

        original_persist = sf._persist_generated_slide

        def racing_persist(db, **kwargs):
            # Simulate another worker having inserted the identical row first,
            # then run the real upsert so the IntegrityError path is taken.
            existing = db.query(GeneratedSlide).filter(GeneratedSlide.render_hash == kwargs["render_hash"]).first()
            if existing is None:
                db.add(
                    GeneratedSlide(
                        slide_key="generated-other",
                        claim_hash=kwargs["claim_hash"],
                        render_hash=kwargs["render_hash"],
                        template_id=kwargs["spec"].template_id,
                        tone=kwargs["spec"].tone,
                        file_key="other-key",
                        status="ready",
                    )
                )
                db.commit()
            return original_persist(db, **kwargs)

        with mock.patch(f"{SLIDE_MODULE}.save_generated_slide_image", return_value="k"), mock.patch(
            f"{SLIDE_MODULE}.file_exists", return_value=False
        ), mock.patch.object(sf, "_persist_generated_slide", side_effect=racing_persist):
            result = realize_generated_slides(
                self.db, plan_a, facts=[], passages=[], fill_fn=fill, vision_fn=vision,
                renderer=FakeRenderer(), brand_page_fn=lambda p: b"x",
            )
        self.assertIsInstance(result[1], GeneratedSlideInstance)
        # Both writers converge on the single row created first.
        self.assertEqual(self.db.query(GeneratedSlide).count(), 1)
        self.assertEqual(result[1].slide_key, "generated-other")
        _ = plan_b


class AttemptAuditTests(MemoryDbTestCase):
    def setUp(self):
        super().setUp()
        user = User(email="attempts@example.com", password_hash="x")
        self.db.add(user)
        self.db.flush()
        generation = Generation(
            user_id=user.id,
            audience_cluster="A1",
            duration="T1",
            channel="live",
            intent="I1",
            temperature="T1",
        )
        self.db.add(generation)
        self.db.commit()
        self.generation_id = generation.id

    def _run(self, *, fill_fn, vision_fn, renderer=None):
        plan = [
            BrandSlide(COVER_PAGE, "", "Cover"),
            placeholder(),
            BrandSlide(CLOSING_PAGE, "M14", "Close"),
        ]
        with mock.patch(
            f"{SLIDE_MODULE}.save_generated_slide_attempt_image",
            return_value=f"generated-slide-attempts/{self.generation_id}/attempt.jpg",
        ) as save_attempt, mock.patch(
            f"{SLIDE_MODULE}.save_generated_slide_image",
            return_value="generated-slides/canonical.jpg",
        ), mock.patch(f"{SLIDE_MODULE}.file_exists", return_value=True):
            result = realize_generated_slides(
                self.db,
                plan,
                generation_id=self.generation_id,
                facts=[],
                passages=[],
                fill_fn=fill_fn,
                vision_fn=vision_fn,
                renderer=renderer or FakeRenderer(),
                brand_page_fn=lambda page: make_jpeg((10, 10, 10)),
            )
        return result, save_attempt

    def test_success_archives_rendered_attempt_and_links_canonical_slide(self):
        result, save_attempt = self._run(
            fill_fn=mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""}),
            vision_fn=mock.Mock(return_value={"passed": True, "violations": []}),
        )

        self.assertIsInstance(result[1], GeneratedSlideInstance)
        attempt = self.db.query(GeneratedSlideAttempt).one()
        self.assertEqual(attempt.generation_id, self.generation_id)
        self.assertEqual(attempt.placeholder_key, placeholder().slide_key)
        self.assertEqual(attempt.attempt_number, 1)
        self.assertEqual(attempt.outcome, "success")
        self.assertEqual(attempt.gate, "accepted")
        self.assertEqual(attempt.review_status, "pending")
        self.assertFalse(attempt.use_as_guidance)
        self.assertEqual(attempt.generated_slide_id, result[1].generated_slide_id)
        self.assertEqual(json.loads(attempt.slot_values_json)["title"], "Not a *bootcamp*")
        self.assertEqual(attempt.file_key, f"generated-slide-attempts/{self.generation_id}/attempt.jpg")
        save_attempt.assert_called_once()
        self.assertEqual(save_attempt.call_args.args[0], self.generation_id)

    def test_failed_attempts_and_final_degradation_are_archived(self):
        result, save_attempt = self._run(
            fill_fn=mock.Mock(return_value={"title": "We place 999 students", "subtitle": ""}),
            vision_fn=mock.Mock(return_value={"passed": True, "violations": []}),
        )

        self.assertIsInstance(result[1], BrandSlide)
        rows = (
            self.db.query(GeneratedSlideAttempt)
            .order_by(GeneratedSlideAttempt.id)
            .all()
        )
        self.assertEqual([row.outcome for row in rows], [
            "provenance",
            "provenance",
            "provenance",
            "final_degradation",
        ])
        self.assertEqual([row.attempt_number for row in rows], [1, 2, 3, 4])
        self.assertTrue(json.loads(rows[0].violations_json))
        save_attempt.assert_not_called()

    def test_furniture_rejection_is_archived_and_suppresses_same_template(self):
        first = placeholder(slide_key="generated:first", claim="First unsupported beat")
        second = placeholder(slide_key="generated:second", claim="Second unsupported beat")
        plan = [
            BrandSlide(COVER_PAGE, "", "Cover"),
            first,
            second,
            BrandSlide(CLOSING_PAGE, "M14", "Close"),
        ]
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""})
        vision = mock.Mock(return_value={
            "passed": False,
            "violations": [{"category": "furniture", "detail": "Brand lockup is missing"}],
        })
        renderer = FakeRenderer()

        with mock.patch(
            f"{SLIDE_MODULE}.save_generated_slide_attempt_image",
            return_value=f"generated-slide-attempts/{self.generation_id}/attempt.jpg",
        ), mock.patch(f"{SLIDE_MODULE}.file_exists", return_value=False):
            result = realize_generated_slides(
                self.db,
                plan,
                generation_id=self.generation_id,
                facts=[],
                passages=[],
                fill_fn=fill,
                vision_fn=vision,
                renderer=renderer,
                brand_page_fn=lambda page: make_jpeg((10, 10, 10)),
            )

        self.assertIsInstance(result[1], BrandSlide)
        self.assertIsInstance(result[2], BrandSlide)
        fill.assert_called_once()
        vision.assert_called_once()
        self.assertEqual(len(renderer.calls), 1)
        rows = (
            self.db.query(GeneratedSlideAttempt)
            .order_by(GeneratedSlideAttempt.id)
            .all()
        )
        self.assertEqual(
            [(row.placeholder_key, row.outcome, row.gate) for row in rows],
            [
                ("generated:first", "vision_furniture", "vision"),
                ("generated:first", "final_degradation", "furniture"),
                ("generated:second", "template_suppressed", "furniture"),
                ("generated:second", "final_degradation", "furniture"),
            ],
        )
        self.assertIn("Brand lockup is missing", rows[0].violations_json)
        self.assertIn("suppressed", rows[2].violations_json.lower())

    def test_only_approved_template_scoped_notes_are_injected_and_deduped(self):
        notes = [
            GeneratedSlideAttempt(
                generation_id=self.generation_id,
                placeholder_key=f"prior-{index}",
                attempt_number=1,
                claim="Prior",
                template_id=template_id,
                tone="light",
                outcome="success",
                gate="accepted",
                review_status=status,
                review_note=note,
                use_as_guidance=use,
            )
            for index, (template_id, status, use, note) in enumerate([
                ("section-divider-light", "approved", True, "Keep the title concrete"),
                ("section-divider-light", "approved", True, "  keep   the title concrete  "),
                ("section-divider-light", "pending", True, "Pending note"),
                ("section-divider-light", "approved", False, "Not opted in"),
                ("stat-light", "approved", True, "Wrong template"),
            ])
        ]
        self.db.add_all(notes)
        self.db.commit()
        fill = mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""})

        self._run(
            fill_fn=fill,
            vision_fn=mock.Mock(return_value={"passed": True, "violations": []}),
        )

        self.assertEqual(
            fill.call_args.args[0]["approved_review_guidance"],
            ["keep the title concrete"],
        )

    def test_audit_content_is_immutable_but_review_fields_are_editable(self):
        self._run(
            fill_fn=mock.Mock(return_value={"title": "Not a *bootcamp*", "subtitle": ""}),
            vision_fn=mock.Mock(return_value={"passed": True, "violations": []}),
        )
        attempt = self.db.query(GeneratedSlideAttempt).one()
        attempt.claim = "Changed claim"
        with self.assertRaisesRegex(ValueError, "immutable"):
            self.db.commit()
        self.db.rollback()

        attempt = self.db.query(GeneratedSlideAttempt).one()
        attempt.review_status = "approved"
        attempt.review_note = "Use this pattern"
        attempt.use_as_guidance = True
        self.db.commit()
        self.assertEqual(attempt.review_status, "approved")


# ---------------------------------------------------------------------------
# Deck spec + PPTX embedding of a concrete generated slide
# ---------------------------------------------------------------------------


class DeckSpecAndPptxTests(unittest.TestCase):
    def _instance(self) -> GeneratedSlideInstance:
        return GeneratedSlideInstance(
            slide_key="generated-abc123",
            template_id="section-divider-light",
            template_version="1",
            tone="light",
            title="Not a bootcamp",
            module_id="M13",
            file_key="generated-slides/abc.jpg",
            generated_slide_id=7,
            claim_hash="c" * 64,
            render_hash="d" * 64,
        )

    def test_deck_slide_carries_generated_source_key_and_url(self):
        deck_slide = deck_slide_from_planned(self._instance())
        self.assertEqual(deck_slide.source, GENERATED_SOURCE)
        self.assertEqual(deck_slide.slide_key, "generated-abc123")
        self.assertEqual(deck_slide.image_url, "/api/generated-slides/generated-abc123.jpg")
        self.assertEqual(deck_slide.page, 0)  # generated slides have no brand page
        self.assertEqual(deck_slide.module_id, "M13")

    def test_render_pptx_embeds_generated_and_brand_slides_from_storage(self):
        plan = [
            BrandSlide(COVER_PAGE, "", "Cover"),
            self._instance(),
            BrandSlide(CLOSING_PAGE, "M14", "Close"),
        ]
        saved = {}

        def fake_save_file(key, data, content_type):
            saved["key"] = key
            saved["data"] = data
            return key

        with mock.patch("backend.pipeline.brand_deck.page_image", return_value=make_jpeg((5, 5, 5))), mock.patch(
            "backend.pipeline.brand_deck.read_file", return_value=make_jpeg((9, 9, 9))
        ) as read_file, mock.patch("backend.pipeline.brand_deck.save_file", side_effect=fake_save_file):
            path = render_pptx(plan, 42, notes_by_slide_key={}, file_key=None)

        self.assertEqual(path, "decks/42.pptx")
        read_file.assert_called_once_with("generated-slides/abc.jpg")

        from pptx import Presentation

        presentation = Presentation(io.BytesIO(saved["data"]))
        self.assertEqual(len(presentation.slides), 3)  # cover + generated + close

    def test_render_pptx_skips_an_unrealised_placeholder(self):
        plan = [
            BrandSlide(COVER_PAGE, "", "Cover"),
            placeholder(),  # never filled: no file_key
            BrandSlide(CLOSING_PAGE, "M14", "Close"),
        ]
        saved = {}

        def fake_save_file(key, data, content_type):
            saved["data"] = data
            return key

        with mock.patch("backend.pipeline.brand_deck.page_image", return_value=make_jpeg()), mock.patch(
            "backend.pipeline.brand_deck.save_file", side_effect=fake_save_file
        ):
            render_pptx(plan, 43, file_key=None)

        from pptx import Presentation

        presentation = Presentation(io.BytesIO(saved["data"]))
        self.assertEqual(len(presentation.slides), 2)  # placeholder skipped, deck not extended


# ---------------------------------------------------------------------------
# Stage 3 metadata: the placeholder remembers the brand page it displaced
# ---------------------------------------------------------------------------


class DisplacedMetadataTests(unittest.TestCase):
    def test_planted_placeholder_records_the_displaced_brand_page(self):
        from backend.pipeline.deck import slide_count_for

        sequence = ["M01", "M04", "M07", "M13", "M14"]
        from backend.pipeline.brand_deck import plan_pages

        plan = plan_pages(sequence, slide_count_for("T2"))
        context = "Here is what we are not: not a glorified weekend bootcamp reselling recycled MOOCs."
        modules = [SimpleNamespace(id="M13", name="The honest boundary", job="We are a real institution", core_content="")]
        result = build_gap_plan(plan, sequence, "T2", context_note=context, modules=modules)

        self.assertEqual(len(result.generated), 1)
        planted = result.generated[0]
        self.assertIsNotNone(planted.replaced_page)
        self.assertIsInstance(planted.replaced_page, int)
        # It is a real brand page that was in the original plan.
        original_pages = {slide.page for slide in plan}
        self.assertIn(planted.replaced_page, original_pages)


if __name__ == "__main__":
    unittest.main()
