from __future__ import annotations

from pathlib import Path
import tempfile
import unittest
from unittest import mock

from backend.config import settings
from backend.pipeline import brand_deck
from backend.pipeline.brand_deck import (
    CLOSING_PAGE,
    COVER_PAGE,
    MODULE_PAGES,
    SOURCE_PAGE_COUNT,
    BrandDeckUnavailable,
    BrandSlide,
    _assign_occurrences,
    brand_slide_key,
    deck_spec_from_plan,
    plan_pages,
    render_pptx,
    resolve_source,
)
from backend.pipeline.deck import (
    BRAND_SOURCE,
    SLIDE_COUNTS,
    PlannedSlide,
    slide_ceiling_for,
    slide_count_for,
)
from backend.pipeline.resolver import FALLBACK_BY_INTENT


class BrandDeckMapTests(unittest.TestCase):
    def test_every_source_page_is_mapped_exactly_once(self):
        pages = [COVER_PAGE, CLOSING_PAGE]
        for entries in MODULE_PAGES.values():
            pages.extend(page for page, _label in entries)
        self.assertEqual(len(pages), len(set(pages)), "a page is mapped to two modules")
        self.assertEqual(sorted(pages), list(range(1, SOURCE_PAGE_COUNT + 1)))

    def test_every_page_has_a_label(self):
        for module_id, entries in MODULE_PAGES.items():
            for page, label in entries:
                self.assertTrue(label.strip(), f"{module_id} p{page} has no label")


class BrandDeckPlanTests(unittest.TestCase):
    def test_ninety_minute_deck_is_longer_than_thirty(self):
        self.assertGreater(SLIDE_COUNTS["T5"], SLIDE_COUNTS["T4"])
        self.assertEqual(slide_count_for("T5"), 50)

    def test_slide_ceiling_for_all_durations(self):
        self.assertEqual(slide_ceiling_for("T0"), 8)
        self.assertEqual(slide_ceiling_for("T1"), 12)
        self.assertEqual(slide_ceiling_for("T2"), 16)
        self.assertEqual(slide_ceiling_for("T3"), 22)
        self.assertEqual(slide_ceiling_for("T4"), 51)
        self.assertEqual(slide_ceiling_for("T5"), 153)

    def test_plan_is_bookended_and_within_budget(self):
        plan = plan_pages(["M01", "M04", "M07", "M14"], 12)
        self.assertEqual(len(plan), 12)
        self.assertEqual(plan[0].page, COVER_PAGE)
        self.assertEqual(plan[-1].page, CLOSING_PAGE)

    def test_no_page_repeats(self):
        plan = plan_pages(["M01", "M04", "M07", "M14"], 22)
        pages = [slide.page for slide in plan]
        self.assertEqual(len(pages), len(set(pages)))

    def test_modules_appear_in_recipe_order(self):
        sequence = ["M07", "M01", "M12"]
        plan = plan_pages(sequence, 16)
        seen: list[str] = []
        for slide in plan[1:-1]:
            if slide.module_id and (not seen or seen[-1] != slide.module_id):
                seen.append(slide.module_id)
        self.assertEqual([mid for mid in seen if mid in sequence][: len(sequence)], sequence)

    def test_pages_within_a_module_stay_in_deck_order(self):
        plan = plan_pages(["M08"], 8)
        pages = [slide.page for slide in plan if slide.module_id == "M08"]
        self.assertEqual(pages, sorted(pages))

    def test_short_deck_takes_each_modules_best_page(self):
        # T0 is 8 slides: cover + closing leaves 6 for three modules.
        plan = plan_pages(["M01", "M04", "M14"], 8)
        self.assertEqual(len(plan), 8)
        self.assertIn(6, [slide.page for slide in plan])  # M01's best page
        self.assertIn(51, [slide.page for slide in plan])  # M04's best page

    def test_thin_sequence_stays_within_ceiling_without_off_recipe_pages(self):
        # M13 has two pages and M14 one — the deck stays short rather than
        # padding with unrelated modules.
        sequence = ["M13", "M14"]
        plan = plan_pages(sequence, 16)
        self.assertLessEqual(len(plan), 16)
        self.assertEqual(plan[0].page, COVER_PAGE)
        self.assertEqual(plan[-1].page, CLOSING_PAGE)
        for slide in plan[1:-1]:
            self.assertIn(slide.module_id, sequence)

    def test_every_recipe_fallback_stays_within_ceiling(self):
        for intent, sequence in FALLBACK_BY_INTENT.items():
            for duration in ("T0", "T1", "T2", "T3", "T4", "T5"):
                ceiling = slide_ceiling_for(duration)
                plan = plan_pages(list(sequence), ceiling)
                self.assertLessEqual(
                    len(plan),
                    ceiling,
                    f"{intent}/{duration} produced {len(plan)} over ceiling {ceiling}",
                )
                for slide in plan[1:-1]:
                    self.assertIn(
                        slide.module_id,
                        sequence,
                        f"{intent}/{duration} body page {slide.page} is off-recipe",
                    )

    def test_t4_plan_includes_every_curated_page_of_sequence(self):
        sequence = ["M01", "M02", "M04", "M07", "M13", "M14"]
        ceiling = slide_ceiling_for("T4")
        plan = plan_pages(sequence, ceiling)
        self.assertLess(len(plan), ceiling)
        planned_pages = {slide.page for slide in plan}
        for module_id in sequence:
            for page, _label in MODULE_PAGES[module_id]:
                self.assertIn(page, planned_pages, f"{module_id} p{page} missing from T4 plan")

    def test_unknown_modules_are_ignored(self):
        plan = plan_pages(["M99", "M01"], 6)
        self.assertEqual(len(plan), 6)
        self.assertTrue(all(slide.module_id != "M99" for slide in plan))

    def test_spec_carries_page_and_image_url(self):
        spec = deck_spec_from_plan(plan_pages(["M01", "M14"], 6))
        self.assertTrue(spec.slides)
        for slide in spec.slides:
            self.assertEqual(slide.layout, "image")
            self.assertGreater(slide.page, 0)
            self.assertEqual(slide.image_url, f"/api/brand-deck/pages/{slide.page}.jpg")


class SlideIdentityTests(unittest.TestCase):
    def test_brand_slide_satisfies_planned_slide_interface(self):
        slide = BrandSlide(6, "M01", "Origin")
        self.assertIsInstance(slide, PlannedSlide)
        self.assertEqual(slide.source, BRAND_SOURCE)
        self.assertEqual(slide.title, "Origin")
        self.assertEqual(slide.page, 6)
        self.assertEqual(slide.slide_key, "brand:p6")

    def test_brand_slide_key_disambiguates_repeated_pages(self):
        self.assertEqual(brand_slide_key(8), "brand:p8")
        self.assertEqual(brand_slide_key(8, 1), "brand:p8")
        self.assertEqual(brand_slide_key(8, 2), "brand:p8#2")
        self.assertEqual(brand_slide_key(8, 3), "brand:p8#3")

    def test_assign_occurrences_gives_repeated_pages_unique_keys(self):
        numbered = _assign_occurrences(
            [
                BrandSlide(1, "", "Cover"),
                BrandSlide(8, "M09", "Gurugram"),
                BrandSlide(20, "M03", "Model"),
                BrandSlide(8, "M09", "Gurugram again"),
            ]
        )
        keys = [slide.slide_key for slide in numbered]
        self.assertEqual(keys, ["brand:p1", "brand:p8", "brand:p20", "brand:p8#2"])
        self.assertEqual(len(keys), len(set(keys)))

    def test_plan_slide_keys_are_unique_and_page_aligned(self):
        plan = plan_pages(["M01", "M04", "M07", "M14"], 22)
        keys = [slide.slide_key for slide in plan]
        self.assertEqual(len(keys), len(set(keys)))
        # No repeats today, so every key is the plain page key.
        self.assertEqual(keys, [f"brand:p{slide.page}" for slide in plan])

    def test_every_recipe_fallback_produces_stable_unique_slide_keys(self):
        # Golden regression: identity is deterministic and collision-free for
        # every recipe fallback across every duration.
        for intent, sequence in FALLBACK_BY_INTENT.items():
            for duration in ("T0", "T1", "T2", "T3", "T4", "T5"):
                plan = plan_pages(list(sequence), slide_ceiling_for(duration))
                keys = [slide.slide_key for slide in plan]
                self.assertEqual(
                    len(keys), len(set(keys)), f"{intent}/{duration} has duplicate slide keys"
                )
                self.assertEqual(
                    keys,
                    [brand_slide_key(slide.page, slide.occurrence) for slide in plan],
                    f"{intent}/{duration} key drifted from its page/occurrence",
                )

    def test_spec_carries_identity_fields(self):
        spec = deck_spec_from_plan(plan_pages(["M01", "M14"], 6))
        for deck_slide, planned in zip(spec.slides, plan_pages(["M01", "M14"], 6)):
            self.assertEqual(deck_slide.source, BRAND_SOURCE)
            self.assertEqual(deck_slide.kind, "image")
            self.assertEqual(deck_slide.slide_key, planned.slide_key)
            self.assertEqual(deck_slide.title, planned.label)

    def test_deck_slide_defaults_keep_old_specs_loadable(self):
        from backend.pipeline.deck import spec_from_dict

        # A deck stored before Stage 0 had no source/kind/slide_key.
        spec = spec_from_dict(
            {"slides": [{"layout": "image", "page": 6, "title": "Origin", "module_id": "M01"}]}
        )
        self.assertEqual(spec.slides[0].source, BRAND_SOURCE)
        self.assertEqual(spec.slides[0].kind, "image")
        self.assertEqual(spec.slides[0].slide_key, "")


class BrandDeckSourceTests(unittest.TestCase):
    def test_cached_source_is_used_without_a_database_file_key(self):
        with tempfile.TemporaryDirectory() as directory:
            files_dir = Path(directory)
            cached = files_dir / brand_deck.SOURCE_CACHE_PATH
            cached.parent.mkdir(parents=True)
            cached.write_bytes(b"cached brand deck")
            with mock.patch.object(brand_deck, "FILES_DIR", files_dir), mock.patch.object(
                brand_deck, "_local_source_candidates", return_value=[]
            ):
                self.assertEqual(brand_deck.resolve_source(None), cached)


class BrandDeckRenderTests(unittest.TestCase):
    """Needs the brand deck PDF, so it skips where the source is not synced."""

    def setUp(self):
        try:
            resolve_source()
        except BrandDeckUnavailable as exc:
            self.skipTest(str(exc))

    def test_renders_one_full_bleed_picture_per_planned_page(self):
        from pptx import Presentation

        plan = plan_pages(["M01", "M07", "M14"], 8)
        with mock.patch.object(type(settings), "uses_spaces", property(lambda self: False)):
            path = render_pptx(plan, 999200, notes_by_module={"M01": "Origin note."})
        presentation = Presentation(path)
        self.assertEqual(len(presentation.slides._sldIdLst), len(plan))
        for slide in presentation.slides:
            picture = slide.shapes[0]
            self.assertEqual(picture.width, presentation.slide_width)
            self.assertEqual(picture.height, presentation.slide_height)

    def test_speaker_notes_follow_selected_pages(self):
        from pptx import Presentation

        plan = plan_pages(["M01", "M07", "M14"], 8)
        page_notes = {slide.page: f"Spoken note for p{slide.page}." for slide in plan}
        with mock.patch.object(type(settings), "uses_spaces", property(lambda self: False)):
            path = render_pptx(
                plan,
                999201,
                notes_by_module={"M01": "Module-wide note that should lose."},
                notes_by_page=page_notes,
            )
        presentation = Presentation(path)
        for slide, item in zip(presentation.slides, plan):
            self.assertEqual(slide.notes_slide.notes_text_frame.text, f"Spoken note for p{item.page}.")

    def test_speaker_notes_follow_slide_keys_over_pages(self):
        from pptx import Presentation

        plan = plan_pages(["M01", "M07", "M14"], 8)
        key_notes = {slide.slide_key: f"Keyed note for {slide.slide_key}." for slide in plan}
        with mock.patch.object(type(settings), "uses_spaces", property(lambda self: False)):
            path = render_pptx(
                plan,
                999202,
                notes_by_page={slide.page: "Page note that should lose." for slide in plan},
                notes_by_slide_key=key_notes,
            )
        presentation = Presentation(path)
        for slide, item in zip(presentation.slides, plan):
            self.assertEqual(
                slide.notes_slide.notes_text_frame.text, f"Keyed note for {item.slide_key}."
            )

    def test_repeated_page_keeps_a_distinct_note_per_occurrence(self):
        from pptx import Presentation

        plan = _assign_occurrences(
            [
                BrandSlide(COVER_PAGE, "", "Learn by Doing"),
                BrandSlide(8, "M09", "Gurugram"),
                BrandSlide(8, "M09", "Gurugram, revisited"),
                BrandSlide(CLOSING_PAGE, "M14", "Close"),
            ]
        )
        key_notes = {
            "brand:p8": "First time we see Gurugram.",
            "brand:p8#2": "We circle back to Gurugram.",
        }
        with mock.patch.object(type(settings), "uses_spaces", property(lambda self: False)):
            path = render_pptx(plan, 999203, notes_by_slide_key=key_notes)
        presentation = Presentation(path)
        notes = [slide.notes_slide.notes_text_frame.text for slide in presentation.slides]
        self.assertIn("First time we see Gurugram.", notes)
        self.assertIn("We circle back to Gurugram.", notes)
