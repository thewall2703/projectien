"""The DS & AI deck (MU x IITian's Hub): a second, opt-in slide source.

The brand deck stays the backbone of every pitch. This deck is only planned in
when the request's free-text context is about data science / AI, and then it
contributes a bounded share of the slide budget *alongside* the brand deck —
never instead of it. Pages come from a PDF export of the source PPTX and are
rasterised and cached exactly like brand deck pages.

Which DS & AI pages a pitch gets is decided by the deck's own topic index
(Brand Deck Testing → DS & AI deck): persona recommendations and reviewer
verdicts on each topic, not a hand-curated page map.
"""

from __future__ import annotations

import json
import re
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from threading import Lock
from typing import Any

from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from backend.config import FILES_DIR
from backend.models import Asset, DeckTopic, utc_now
from backend.storage import read_file, save_file

DSAI_SOURCE = "dsai"
DECK_KEY = "dsai"
DECK_NAME = "DS & AI Deck"
ASSET_TYPE = "deck"
ASSET_TITLE = "DS & AI Deck (MU x IITian's Hub)"
STORAGE_PDF_KEY = "library/decks/dsai-deck.pdf"
SOURCE_CACHE_PATH = "dsai-deck/source.pdf"
PAGE_CACHE_PREFIX = "dsai-deck/pages"
IMAGE_DPI = 72
JPEG_QUALITY = 82
MAX_LABEL_CHARS = 70

_DSAI_PATTERN = re.compile(
    r"\b("
    r"ds\s*[/&-]?\s*ai|dsai|data\s+scien\w*|artificial\s+intelligence|machine\s+learning"
    r"|gen\s*ai|ai\s*/\s*ml|ai|ml|emerging\s+tech\w*"
    r")\b",
    re.IGNORECASE,
)


class DsaiDeckUnavailable(RuntimeError):
    """Raised when the DS & AI deck PDF cannot be located."""


def is_dsai_request(*texts: str | None) -> bool:
    """True when the generation request's free text asks for the DS & AI use case."""
    return any(_DSAI_PATTERN.search(text or "") for text in texts)


def dsai_slide_budget(ceiling: int) -> int:
    """Slides reserved for the DS & AI deck; the brand deck keeps the rest."""
    return max(2, ceiling // 3)


def dsai_slide_key(page: int) -> str:
    return f"{DSAI_SOURCE}:p{page}"


def dsai_image_url(page: int) -> str:
    return f"/api/dsai-deck/pages/{page}.jpg"


@dataclass(frozen=True)
class DsaiSlide:
    """One DS & AI deck page placed in a generated deck (a ``PlannedSlide``)."""

    page: int
    module_id: str
    label: str

    @property
    def source(self) -> str:
        return DSAI_SOURCE

    @property
    def title(self) -> str:
        return self.label

    @property
    def slide_key(self) -> str:
        return dsai_slide_key(self.page)

    @property
    def image_url(self) -> str:
        return dsai_image_url(self.page)


# ---------------------------------------------------------------------------
# Asset, labels and page images
# ---------------------------------------------------------------------------

_labels_cache: dict[int, str] = {}


def find_dsai_asset(db: Session) -> Asset | None:
    try:
        asset = (
            db.query(Asset)
            .filter(Asset.type == ASSET_TYPE, Asset.title == ASSET_TITLE)
            .order_by(Asset.id)
            .first()
        )
    except SQLAlchemyError:
        return None
    if asset is not None:
        _labels_cache.clear()
        _labels_cache.update(page_labels(asset))
    return asset


def page_entries(asset: Asset | None) -> list[dict[str, Any]]:
    if asset is None or not asset.extract_json:
        return []
    try:
        payload = json.loads(asset.extract_json)
    except json.JSONDecodeError:
        return []
    return [item for item in payload.get("pages") or [] if isinstance(item, dict)]


def page_count(asset: Asset | None) -> int:
    return len(page_entries(asset))


def page_labels(asset: Asset | None) -> dict[int, str]:
    labels: dict[int, str] = {}
    for item in page_entries(asset):
        try:
            page = int(item.get("page") or 0)
        except (TypeError, ValueError):
            continue
        if page > 0:
            labels[page] = str(item.get("label") or f"DS & AI slide {page}")
    return labels


def cached_label(page: int) -> str:
    return _labels_cache.get(page, f"DS & AI slide {page}")


def resolve_source(file_key: str | None = None) -> Path:
    cached = FILES_DIR / SOURCE_CACHE_PATH
    if cached.exists():
        return cached
    try:
        data = read_file(file_key or STORAGE_PDF_KEY)
    except Exception as exc:  # noqa: BLE001
        raise DsaiDeckUnavailable(
            "DS & AI deck PDF not found. Import it with `python -m backend.pipeline.dsai_deck <pptx>`."
        ) from exc
    cached.parent.mkdir(parents=True, exist_ok=True)
    cached.write_bytes(data)
    return cached


_source_lock = Lock()
_source_cache: tuple[str, object] | None = None


def _open_source(file_key: str | None = None):
    global _source_cache
    import pymupdf

    path = str(resolve_source(file_key))
    with _source_lock:
        if _source_cache is None or _source_cache[0] != path:
            _source_cache = (path, pymupdf.open(path))
        return _source_cache[1]


def page_image(page: int, file_key: str | None = None) -> bytes:
    cached = FILES_DIR / PAGE_CACHE_PREFIX / f"p{page:03d}.jpg"
    if cached.exists():
        return cached.read_bytes()
    document = _open_source(file_key)
    if page < 1 or page > len(document):
        raise ValueError(f"page {page} is outside the DS & AI deck")
    data = document[page - 1].get_pixmap(dpi=IMAGE_DPI).tobytes("jpeg", jpg_quality=JPEG_QUALITY)
    cached.parent.mkdir(parents=True, exist_ok=True)
    cached.write_bytes(data)
    return data


# ---------------------------------------------------------------------------
# Import: PPTX -> PDF + per-page text
# ---------------------------------------------------------------------------


def _slide_texts(pptx_path: Path) -> list[str]:
    from pptx import Presentation

    texts: list[str] = []
    for slide in Presentation(str(pptx_path)).slides:
        parts = [
            " ".join(shape.text_frame.text.split())
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        texts.append(" · ".join(part for part in parts if part))
    return texts


def _label_for(page: int, text: str) -> str:
    text = text.strip()
    if not text:
        return f"DS & AI slide {page}"
    if len(text) > MAX_LABEL_CHARS:
        text = text[:MAX_LABEL_CHARS].rsplit(" ", 1)[0] + "…"
    return text


def convert_to_pdf(pptx_path: Path) -> bytes:
    with tempfile.TemporaryDirectory() as out_dir:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, str(pptx_path)],
            check=True,
            capture_output=True,
            timeout=900,
        )
        pdfs = list(Path(out_dir).glob("*.pdf"))
        if not pdfs:
            raise DsaiDeckUnavailable("LibreOffice did not produce a PDF")
        return pdfs[0].read_bytes()


def import_deck(db: Session, pptx_path: Path, pdf_bytes: bytes | None = None) -> Asset:
    """Store the deck PDF and page text, and upsert its library asset."""
    import pymupdf

    texts = _slide_texts(pptx_path)
    pdf = pdf_bytes if pdf_bytes is not None else convert_to_pdf(pptx_path)
    pdf_pages = len(pymupdf.open(stream=pdf, filetype="pdf"))
    if pdf_pages != len(texts):
        raise DsaiDeckUnavailable(f"PDF has {pdf_pages} pages but the PPTX has {len(texts)} slides")
    file_key = save_file(STORAGE_PDF_KEY, pdf, "application/pdf")

    cached = FILES_DIR / SOURCE_CACHE_PATH
    cached.parent.mkdir(parents=True, exist_ok=True)
    cached.write_bytes(pdf)
    for stale in (FILES_DIR / PAGE_CACHE_PREFIX).glob("p*.jpg"):
        stale.unlink()

    pages = [
        {"page": index, "label": _label_for(index, text), "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    asset = find_dsai_asset(db)
    if asset is None:
        asset = Asset(type=ASSET_TYPE, title=ASSET_TITLE)
        db.add(asset)
    asset.file_key = file_key
    asset.content_type = "application/pdf"
    asset.file_status = "stored"
    asset.status = "exists"
    asset.extract_status = "ready"
    asset.extract_json = json.dumps({"pages": pages}, ensure_ascii=False)
    asset.extracted_at = utc_now()
    asset.edited = True
    asset.notes = "Only used when a generation request's context mentions DS / AI."
    db.commit()
    db.refresh(asset)
    find_dsai_asset(db)
    return asset


# ---------------------------------------------------------------------------
# Planning
# ---------------------------------------------------------------------------


def _topic_rank(row: Any, recipe_ref: str) -> tuple[float, float] | None:
    from backend.media_index import parse_feedback, parse_recommendations

    feedback = parse_feedback(getattr(row, "feedback_json", "") or "")
    verdict = (feedback["verdicts"].get(recipe_ref) or {}).get("verdict") if recipe_ref else None
    if verdict == "no":
        return None
    items = [item for item in parse_recommendations(getattr(row, "recommendations_json", "") or "")["items"] if isinstance(item, dict)]
    score = 0.0
    if verdict == "yes" or any(entry.get("recipe_ref") == recipe_ref for entry in feedback["added"] if isinstance(entry, dict)):
        score += 3.0
    for item in items:
        if recipe_ref and item.get("recipe_ref") == recipe_ref:
            score += 1.0 + float(item.get("confidence") or 0)
    best_any = max((float(item.get("confidence") or 0) for item in items), default=0.0)
    if score <= 0 and best_any <= 0:
        return None
    return score, best_any


def select_dsai_slides(
    topics: list[Any],
    *,
    recipe_ref: str,
    sequence: list[str],
    budget: int,
    labels: dict[int, str] | None = None,
) -> list[DsaiSlide]:
    """Pick up to ``budget`` DS & AI pages for this persona, in deck order.

    Topics the reviewer rejected for the persona are skipped, and topics nobody
    recommended (e.g. the classroom exercises) are never used. The rest are
    ranked persona-first and pages are dealt round-robin, best topic first.
    """
    from backend.deck_topic_index import parse_pages

    if budget <= 0:
        return []
    ranked: list[tuple[tuple[float, float], int, Any]] = []
    for row in topics:
        rank = _topic_rank(row, recipe_ref)
        if rank is not None and parse_pages(getattr(row, "pages_json", "") or ""):
            ranked.append((rank, int(getattr(row, "sort_order", 0) or 0), row))
    ranked.sort(key=lambda item: (-item[0][0], -item[0][1], item[1]))
    pools = [(row, list(parse_pages(row.pages_json))) for _rank, _order, row in ranked]

    sequence_set = set(sequence)
    chosen: dict[int, str] = {}
    while len(chosen) < budget and any(pages for _row, pages in pools):
        for row, pages in pools:
            if len(chosen) >= budget:
                break
            if not pages:
                continue
            page = pages.pop(0)
            modules = [part.strip() for part in (row.module_ids or "").split(",") if part.strip()]
            chosen[page] = next((module for module in modules if module in sequence_set), "")
    names = labels or _labels_cache
    return [
        DsaiSlide(page, module_id, names.get(page, f"DS & AI slide {page}"))
        for page, module_id in sorted(chosen.items())
    ]


def insert_before_closing(plan: list[Any], slides: list[DsaiSlide]) -> list[Any]:
    """Place the DS & AI section just before the brand deck's closing slide."""
    from backend.pipeline.brand_deck import CLOSING_PAGE
    from backend.pipeline.deck import BRAND_SOURCE

    if not slides:
        return list(plan)
    index = len(plan)
    for position in range(len(plan) - 1, -1, -1):
        slide = plan[position]
        if getattr(slide, "source", "") == BRAND_SOURCE and getattr(slide, "page", None) == CLOSING_PAGE:
            index = position
            break
    return [*plan[:index], *slides, *plan[index:]]


def load_dsai_slides(
    db: Session,
    *,
    recipe_ref: str,
    sequence: list[str],
    budget: int,
) -> list[DsaiSlide]:
    asset = find_dsai_asset(db)
    if asset is None or budget <= 0:
        return []
    topics = (
        db.query(DeckTopic)
        .filter(DeckTopic.deck == DECK_KEY)
        .order_by(DeckTopic.sort_order, DeckTopic.id)
        .all()
    )
    return select_dsai_slides(
        topics,
        recipe_ref=recipe_ref,
        sequence=sequence,
        budget=budget,
        labels=page_labels(asset),
    )


if __name__ == "__main__":
    import sys

    from backend.database import SessionLocal, ensure_schema

    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m backend.pipeline.dsai_deck <deck.pptx>")
    ensure_schema()
    session = SessionLocal()
    try:
        imported = import_deck(session, Path(sys.argv[1]))
        print(f"Imported {page_count(imported)} pages into asset {imported.id} ({imported.file_key})")
    finally:
        session.close()
