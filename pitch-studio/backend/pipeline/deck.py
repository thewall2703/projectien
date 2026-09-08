from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pydantic import BaseModel, Field

from backend.storage import save_file

INK = RGBColor(0x14, 0x14, 0x1A)
PAPER = RGBColor(0xF5, 0xF2, 0xEB)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x24, 0x39, 0xD9)
ACCENT_DARK = RGBColor(0x16, 0x22, 0x7A)
GOLD = RGBColor(0xC9, 0xA2, 0x4B)
MUTED = RGBColor(0x6B, 0x6B, 0x76)
LINE = RGBColor(0xE7, 0xE2, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = 13.333
SLIDE_H = 7.5

SLIDE_COUNTS = {
    "T0": 8,
    "T1": 12,
    "T2": 16,
    "T3": 22,
    "T4": 30,
    "T5": 26,
}

_SENTENCE_SPLIT = re.compile(r"(?<=[.!?])\s+")
_STAT_VALUE = re.compile(
    r"(₹\s?[\d,.]+(?:\s*LPA)?|[\d]{1,2}(?:\.\d+)?\s*%|[\d]{1,3}(?:,\d{3})+(?:\.\d+)?|\b\d{2,4}\b)"
)

LAYOUTS = {"title", "section", "stat_pair", "bullets", "profile", "quote", "inventory", "cta", "agenda"}
CHROME_FREE = {"title", "section", "cta"}
MODULE_NAMES = {
    "M01": "Origin",
    "M02": "Institution",
    "M03": "Model",
    "M04": "Sequence",
    "M05": "People proof",
    "M06": "Venture proof",
    "M07": "Outcome proof",
    "M08": "Faculty",
    "M09": "Campus",
    "M10": "Programmes",
    "M11": "Ecosystem",
    "M12": "Culture",
    "M13": "Honest",
    "M14": "Ask",
}


def _module_name(module_id: str) -> str:
    return MODULE_NAMES.get(module_id or "", module_id or "")


class DeckStat(BaseModel):
    label: str = ""
    value: str = ""


class DeckSlide(BaseModel):
    layout: str = "bullets"
    title: str = ""
    subtitle: str | None = None
    bullets: list[str] | None = None
    stats: list[DeckStat] | None = None
    quote: str | None = None
    attribution: str | None = None
    module_id: str = ""


class DeckSpec(BaseModel):
    slides: list[DeckSlide] = Field(default_factory=list)


def slide_count_for(duration: str) -> int:
    return SLIDE_COUNTS.get(duration, 16)


def _sentences(text: str) -> list[str]:
    parts = [part.strip() for part in _SENTENCE_SPLIT.split(text or "") if part.strip()]
    return [part[:220] for part in parts if len(part) > 8]


def _chunk(items: list[str], size: int) -> list[list[str]]:
    return [items[index : index + size] for index in range(0, len(items), size)]


def _stats_from_text(text: str) -> list[DeckStat]:
    stats: list[DeckStat] = []
    seen: set[str] = set()
    for match in _STAT_VALUE.finditer(text or ""):
        value = re.sub(r"\s+", " ", match.group(0)).strip()
        if value in seen or (value.isdigit() and int(value) < 10):
            continue
        seen.add(value)
        start = max(0, match.start() - 40)
        label = re.sub(r"\s+", " ", text[start : match.start()].strip(" ,.;:-"))
        label = " ".join(label.split()[-5:]) or "Figure"
        stats.append(DeckStat(label=label[:48], value=value))
        if len(stats) == 4:
            break
    return stats


def expand_deck_from_script(script: dict[str, Any], sequence: list[str]) -> DeckSpec:
    """Build a full deck from the script so a thin LLM spec never ships."""
    sections = script.get("sections") or []
    by_id = {section.get("module_id"): section for section in sections}
    ordered = [by_id[mid] for mid in sequence if mid in by_id]
    if not ordered:
        ordered = sections
    first = ordered[0] if ordered else {}
    slides: list[DeckSlide] = [
        DeckSlide(
            layout="title",
            title=first.get("heading") or "Masters' Union",
            subtitle="One story. The proof, in order.",
            module_id=first.get("module_id") or (sequence[0] if sequence else ""),
        ),
        DeckSlide(
            layout="agenda",
            title="What we will cover",
            bullets=[section.get("heading") or section.get("module_id", "") for section in ordered],
            module_id=first.get("module_id") or "",
        ),
    ]
    for section in ordered:
        module_id = section.get("module_id") or ""
        heading = section.get("heading") or module_id
        text = section.get("text") or ""
        slides.append(DeckSlide(layout="section", title=heading, module_id=module_id))
        points = _sentences(text)
        if not points:
            points = [text.strip() or heading]
        for group in _chunk(points, 4)[:3]:
            slides.append(
                DeckSlide(
                    layout="bullets",
                    title=heading,
                    bullets=group,
                    module_id=module_id,
                )
            )
        stats = _stats_from_text(text)
        if stats:
            slides.append(
                DeckSlide(
                    layout="stat_pair",
                    title=f"{heading} — in numbers",
                    stats=stats,
                    module_id=module_id,
                )
            )
        quoted = re.findall(r"[“\"]([^”\"]{24,180})[”\"]", text)
        if quoted:
            slides.append(
                DeckSlide(
                    layout="quote",
                    title=heading,
                    quote=quoted[0],
                    attribution="From the approved script",
                    module_id=module_id,
                )
            )
    cta = (script.get("cta") or "Ask for the next conversation").strip()
    last_id = ordered[-1].get("module_id") if ordered else (sequence[-1] if sequence else "")
    slides.append(DeckSlide(layout="cta", title=cta, subtitle="Masters' Union", module_id=last_id or ""))
    return DeckSpec(slides=slides)


def merge_deck_specs(generated: DeckSpec, fallback: DeckSpec, sequence: list[str]) -> DeckSpec:
    """Keep the model deck when it is already rich; otherwise use the script floor."""
    content = [
        slide
        for slide in generated.slides
        if slide.layout not in {"title", "section", "agenda", "cta"}
        and (slide.bullets or slide.stats or slide.quote)
    ]
    needed = max(len(sequence) * 2, 6)
    if len(content) >= needed:
        return normalize_deck_order(generated, sequence)
    return fallback


def normalize_deck_order(spec: DeckSpec, sequence: list[str]) -> DeckSpec:
    """Return a deck whose slides follow the script sequence.

    The deck spec is produced by an LLM, which can occasionally emit slides
    whose ``module_id`` values are out of order relative to the script. Rather
    than failing the whole generation, reorder the slides to match ``sequence``.

    A stable sort is used so slides mapped to the same module keep their
    relative order. Slides with no (or unrecognized) ``module_id`` inherit the
    position of the preceding recognized module, keeping them anchored to their
    section (and leading un-anchored slides, e.g. a title, stay at the front).
    """
    index_by_id = {module_id: position for position, module_id in enumerate(sequence)}
    keys: list[int] = []
    current = -1
    for slide in spec.slides:
        if slide.module_id in index_by_id:
            current = index_by_id[slide.module_id]
        keys.append(current)
    ordered = [slide for _, slide in sorted(zip(keys, spec.slides), key=lambda pair: pair[0])]
    return DeckSpec(slides=ordered)


def _set_run(run, text: str, size: int, bold: bool = False, color=INK, font: str = "Georgia") -> None:
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _fill_slide(slide, color=SURFACE) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_shape(slide, left, top, width, height, fill, line=None, rounded: bool = False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    if rounded:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def _add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size,
    bold=False,
    color=INK,
    font="Georgia",
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    _set_run(run, text, size, bold, color, font)
    return box


def _kicker(slide, text: str, left=0.8, top=0.42, width=11.7) -> None:
    _add_text(
        slide,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(0.32),
        (text or "").upper(),
        11,
        True,
        ACCENT,
        "Arial",
    )


def _rule(slide, left, top, width, color=GOLD, height=0.035) -> None:
    _add_shape(slide, Inches(left), Inches(top), Inches(width), Inches(height), color)


def _card(slide, left, top, width, height, fill=SURFACE):
    return _add_shape(slide, Inches(left), Inches(top), Inches(width), Inches(height), fill, LINE, rounded=True)


def _footer(slide, index: int, total: int, module_id: str) -> None:
    _add_shape(slide, Inches(0.8), Inches(7.08), Inches(11.7), Inches(0.012), LINE)
    _add_text(slide, Inches(0.8), Inches(7.14), Inches(5.4), Inches(0.26), "Masters' Union", 10, False, MUTED, "Arial")
    tag = _module_name(module_id)
    _add_text(
        slide,
        Inches(6.2),
        Inches(7.14),
        Inches(3.2),
        Inches(0.26),
        tag,
        10,
        False,
        MUTED,
        "Arial",
        PP_ALIGN.CENTER,
    )
    _add_text(
        slide,
        Inches(9.6),
        Inches(7.14),
        Inches(2.9),
        Inches(0.26),
        f"{index:02d}  /  {total:02d}",
        10,
        False,
        MUTED,
        "Arial",
        PP_ALIGN.RIGHT,
    )


def _render_title(slide, item: DeckSlide) -> None:
    _add_shape(slide, Inches(0), Inches(0), Inches(SLIDE_W), Inches(0.18), ACCENT)
    _kicker(slide, "Masters' Union", 0.9, 2.05)
    _add_text(slide, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.8), item.title or "Masters' Union", 40, True)
    _rule(slide, 0.9, 4.45, 2.2)
    if item.subtitle:
        _add_text(
            slide,
            Inches(0.9),
            Inches(4.65),
            Inches(11.5),
            Inches(1.2),
            item.subtitle,
            18,
            False,
            MUTED,
            "Arial",
        )


def _render_section(slide, item: DeckSlide, section_index: int = 1) -> None:
    _add_text(
        slide,
        Inches(0.7),
        Inches(1.55),
        Inches(11.8),
        Inches(2.2),
        f"{section_index:02d}",
        96,
        True,
        LINE,
        "Georgia",
    )
    _add_text(
        slide,
        Inches(0.9),
        Inches(3.55),
        Inches(11.5),
        Inches(1.4),
        item.title or item.module_id,
        34,
        True,
        ACCENT_DARK,
    )
    if item.subtitle:
        _add_text(
            slide,
            Inches(0.9),
            Inches(5.05),
            Inches(11.5),
            Inches(0.9),
            item.subtitle,
            16,
            False,
            MUTED,
            "Arial",
        )


def _bullet_box(slide, left, top, width, height, bullets: list[str], color=INK, size=17) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    if not bullets:
        return
    for index, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(12)
        run = p.add_run()
        _set_run(run, f"—  {bullet}", size, False, color, "Arial")


def _render_bullets(slide, item: DeckSlide) -> None:
    _kicker(slide, _module_name(item.module_id) or "Key points")
    _add_text(slide, Inches(0.8), Inches(0.78), Inches(11.7), Inches(0.9), item.title or "Key points", 28, True)
    bullets = (item.bullets or [])[:6]
    if not bullets:
        _add_text(
            slide,
            Inches(0.85),
            Inches(1.85),
            Inches(11.5),
            Inches(4.6),
            item.subtitle or "",
            16,
            False,
            INK,
            "Arial",
        )
        return
    _bullet_box(slide, 0.85, 1.85, 11.5, 4.8, bullets)


def _render_stat_pair(slide, item: DeckSlide) -> None:
    _kicker(slide, _module_name(item.module_id) or "Outcomes")
    _add_text(slide, Inches(0.8), Inches(0.78), Inches(11.7), Inches(0.8), item.title or "Outcomes", 28, True)
    stats = (item.stats or [])[:4] or [DeckStat(label="Fact", value="—")]
    count = len(stats)
    if count <= 2:
        width, height, top = 5.6, 3.4, 2.0
    else:
        width, height, top = 5.6, 2.15, 1.9
    for index, stat in enumerate(stats):
        col = index % 2
        row = 0 if count <= 2 else index // 2
        left = 0.8 + col * (width + 0.35)
        card_top = top + row * (height + 0.25)
        _card(slide, left, card_top, width, height, PAPER)
        _add_text(
            slide,
            Inches(left + 0.28),
            Inches(card_top + 0.35),
            Inches(width - 0.5),
            Inches(1.15),
            stat.value or "—",
            36,
            True,
            ACCENT,
        )
        _add_text(
            slide,
            Inches(left + 0.28),
            Inches(card_top + height - 0.85),
            Inches(width - 0.5),
            Inches(0.55),
            stat.label,
            13,
            False,
            MUTED,
            "Arial",
        )


def _render_profile(slide, item: DeckSlide) -> None:
    _kicker(slide, _module_name(item.module_id) or "People")
    _add_text(slide, Inches(0.8), Inches(0.78), Inches(11.7), Inches(0.7), item.title or "People", 28, True)
    if item.subtitle:
        _add_text(slide, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.4), item.subtitle, 14, False, ACCENT, "Arial")
    bullets = (item.bullets or [])[:6]
    if not bullets:
        return
    cols = 2
    width, height = 5.6, 1.55
    start_top = 2.15
    for index, bullet in enumerate(bullets):
        col = index % cols
        row = index // cols
        left = 0.8 + col * (width + 0.35)
        top = start_top + row * (height + 0.22)
        _card(slide, left, top, width, height, PAPER)
        _add_text(
            slide,
            Inches(left + 0.28),
            Inches(top + 0.38),
            Inches(width - 0.5),
            Inches(0.85),
            bullet,
            15,
            False,
            INK,
            "Arial",
        )


def _render_quote(slide, item: DeckSlide) -> None:
    _add_shape(slide, Inches(0.8), Inches(1.7), Inches(0.12), Inches(3.6), ACCENT)
    quote = item.quote or item.title or ""
    _add_text(slide, Inches(1.3), Inches(1.85), Inches(10.8), Inches(2.8), f"“{quote}”", 28, True)
    if item.attribution:
        _add_text(
            slide,
            Inches(1.3),
            Inches(4.9),
            Inches(10.8),
            Inches(0.5),
            item.attribution,
            14,
            False,
            MUTED,
            "Arial",
        )


def _render_inventory(slide, item: DeckSlide) -> None:
    _kicker(slide, _module_name(item.module_id) or "Inventory")
    _add_text(slide, Inches(0.8), Inches(0.78), Inches(11.7), Inches(0.8), item.title or "Inventory", 28, True)
    bullets = (item.bullets or [])[:8]
    if not bullets:
        _add_text(
            slide,
            Inches(0.85),
            Inches(1.85),
            Inches(11.5),
            Inches(4.6),
            item.subtitle or "",
            16,
            False,
            INK,
            "Arial",
        )
        return
    mid = (len(bullets) + 1) // 2
    _bullet_box(slide, 0.85, 1.85, 5.6, 4.6, bullets[:mid])
    _bullet_box(slide, 6.8, 1.85, 5.6, 4.6, bullets[mid:])


def _render_agenda(slide, item: DeckSlide) -> None:
    _kicker(slide, "Agenda")
    _add_text(slide, Inches(0.8), Inches(0.78), Inches(11.7), Inches(0.8), item.title or "What we will cover", 28, True)
    bullets = (item.bullets or [])[:8]
    if not bullets and item.subtitle:
        bullets = [item.subtitle]
    box = slide.shapes.add_textbox(Inches(0.85), Inches(1.85), Inches(11.5), Inches(4.8))
    tf = box.text_frame
    tf.word_wrap = True
    for index, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        num = p.add_run()
        _set_run(num, f"{index + 1:02d}   ", 16, True, ACCENT, "Arial")
        label = p.add_run()
        _set_run(label, bullet, 18, False, INK, "Georgia")


def _render_cta(slide, item: DeckSlide) -> None:
    _add_shape(slide, Inches(0), Inches(0), Inches(SLIDE_W), Inches(SLIDE_H), ACCENT)
    _add_text(slide, Inches(0.9), Inches(2.15), Inches(11.5), Inches(0.4), "NEXT STEP", 12, True, GOLD, "Arial")
    _add_text(
        slide,
        Inches(0.9),
        Inches(2.65),
        Inches(11.5),
        Inches(1.5),
        item.title or "Next step",
        36,
        True,
        WHITE,
    )
    if item.subtitle:
        _add_text(
            slide,
            Inches(0.9),
            Inches(4.35),
            Inches(11.5),
            Inches(1.1),
            item.subtitle,
            18,
            False,
            WHITE,
            "Arial",
        )


RENDERERS = {
    "title": _render_title,
    "stat_pair": _render_stat_pair,
    "bullets": _render_bullets,
    "profile": _render_profile,
    "quote": _render_quote,
    "inventory": _render_inventory,
    "cta": _render_cta,
    "agenda": _render_agenda,
}


def spec_from_dict(payload: dict[str, Any]) -> DeckSpec:
    return DeckSpec.model_validate(payload)


def _apply_notes(slide, text: str) -> None:
    notes = slide.notes_slide
    frame = notes.notes_text_frame
    frame.text = text


def render_pptx(
    spec: DeckSpec,
    generation_id: int,
    sequence: list[str],
    notes_by_module: dict[str, str] | None = None,
) -> str:
    spec = normalize_deck_order(spec, sequence)
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    blank = presentation.slide_layouts[6]
    slides = spec.slides or [
        DeckSlide(layout="title", title="Masters' Union", module_id=sequence[0] if sequence else "")
    ]
    section_index = 0
    total = len(slides)
    for index, item in enumerate(slides, start=1):
        layout = item.layout if item.layout in LAYOUTS else "bullets"
        slide = presentation.slides.add_slide(blank)
        if layout == "cta":
            _fill_slide(slide, ACCENT)
        elif layout == "section":
            _fill_slide(slide, PAPER)
        else:
            _fill_slide(slide, SURFACE)
        if layout == "section":
            section_index += 1
            _render_section(slide, item, section_index)
        else:
            RENDERERS[layout](slide, item)
        if layout not in CHROME_FREE:
            _footer(slide, index, total, item.module_id)
        note = (notes_by_module or {}).get(item.module_id, "")
        if note:
            _apply_notes(slide, note)
    buffer = BytesIO()
    presentation.save(buffer)
    key = f"decks/{generation_id}.pptx"
    return save_file(
        key,
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
