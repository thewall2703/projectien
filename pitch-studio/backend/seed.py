from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.asset_map import infer_matrix_ref, infer_module_ids
from backend.config import DEFAULT_XLSX
from backend.database import SessionLocal, engine, ensure_schema
from backend.models import Asset, Base, DeckVisionRow, LockedFact, Module, Objection, Recipe
from backend.schemas import SeedCounts

OWNER_ROW_MARKERS = ("devansh", "ayushi")


def _cell(row: tuple, index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _normalize_fact_status(raw: str) -> str:
    text = raw.lower()
    if "do not use" in text or "do_not_use" in text:
        return "do_not_use"
    if "conflict" in text:
        return "conflict"
    if "needs decision" in text or "needs_decision" in text:
        return "needs_decision"
    if "verified" in text:
        return "verified"
    return "needs_source"


def _first_token(value: str) -> str:
    return (value.strip().split() or [""])[0]


def _normalize_cluster(value: str) -> str:
    text = value.strip()
    if not text:
        return ""
    if text.lower() == "all":
        return "All"
    return text[0].upper()


def _sort_order(module_id: str) -> int:
    match = re.search(r"(\d+)", module_id)
    return int(match.group(1)) if match else 0


def _upsert_module(db: Session, payload: dict) -> None:
    existing = db.get(Module, payload["id"])
    if existing and existing.edited:
        return
    if existing:
        for key, value in payload.items():
            setattr(existing, key, value)
        existing.edited = False
    else:
        db.add(Module(**payload, edited=False))


def _upsert_by_key(db: Session, model, key_field: str, key_value, payload: dict) -> None:
    existing = db.query(model).filter(getattr(model, key_field) == key_value).first()
    if existing and getattr(existing, "edited", False):
        return
    if existing:
        for key, value in payload.items():
            setattr(existing, key, value)
        existing.edited = False
    else:
        db.add(model(**payload, edited=False))


def seed_modules(db: Session, rows: list[tuple]) -> int:
    count = 0
    for row in rows[1:]:
        module_id = _cell(row, 0)
        if not module_id.startswith("M"):
            continue
        _upsert_module(
            db,
            {
                "id": module_id,
                "name": _cell(row, 1),
                "job": _cell(row, 2),
                "core_content": _cell(row, 3),
                "flex_points": "",
                "sources": _cell(row, 4),
                "owner": _cell(row, 5),
                "sort_order": _sort_order(module_id),
            },
        )
        count += 1
    return count


def seed_facts(db: Session, rows: list[tuple]) -> int:
    count = 0
    for row in rows[1:]:
        fact = _cell(row, 0)
        if not fact:
            continue
        payload = {
            "fact": fact,
            "value": _cell(row, 1),
            "source": _cell(row, 2),
            "status": _normalize_fact_status(_cell(row, 3)),
            "note": _cell(row, 4),
            "module_ids": "",
        }
        existing = db.query(LockedFact).filter(LockedFact.fact == fact).first()
        if existing and existing.edited:
            continue
        if existing:
            for key, value in payload.items():
                setattr(existing, key, value)
            existing.edited = False
        else:
            db.add(LockedFact(**payload, edited=False))
        count += 1
    return count


def seed_recipes(db: Session, rows: list[tuple]) -> int:
    count = 0
    from backend.pipeline.resolver import WORD_BUDGETS

    for row in rows[1:]:
        ref = _cell(row, 0)
        if not ref or ref.lower() == "totals":
            break
        duration = _first_token(_cell(row, 3))
        sequence = _cell(row, 6).replace(" > ", ">").replace(" ", "")
        payload = {
            "ref": ref,
            "audience_label": _cell(row, 1),
            "audience_cluster": _normalize_cluster(_cell(row, 2)),
            "duration": duration,
            "channel": _first_token(_cell(row, 4)),
            "intent": _first_token(_cell(row, 5)),
            "temperature": None,
            "module_sequence": sequence,
            "word_budget": WORD_BUDGETS.get(duration, 0),
            "priority": _cell(row, 7) or "P2",
            "owner": _cell(row, 8),
        }
        _upsert_by_key(db, Recipe, "ref", ref, payload)
        count += 1
    return count


def seed_objections(db: Session, rows: list[tuple]) -> int:
    count = 0
    for row in rows[1:]:
        question = _cell(row, 1)
        if not question:
            continue
        answer = _cell(row, 4)
        status = (
            "blocking"
            if "BLOCKING" in answer or "Needs a written" in answer
            else "approved"
        )
        payload = {
            "question": question,
            "who_asks": _cell(row, 2),
            "move": _cell(row, 3),
            "answer": answer,
            "status": status,
        }
        existing = db.query(Objection).filter(Objection.question == question).first()
        if existing and existing.edited:
            continue
        if existing:
            for key, value in payload.items():
                setattr(existing, key, value)
            existing.edited = False
        else:
            db.add(Objection(**payload, edited=False))
        count += 1
    return count


def _is_owner_row(cells: list[str]) -> bool:
    joined = " ".join(cells).lower()
    return any(marker in joined for marker in OWNER_ROW_MARKERS) and "+" in joined


def seed_assets(db: Session, rows: list[list[tuple[str, str]]]) -> int:
    count = 0
    types = ("video", "photo", "report")
    seen: dict[str, set[str]] = {asset_type: set() for asset_type in types}
    for cells in rows:
        titles = [title for title, _link in cells]
        if not any(titles):
            continue
        if _is_owner_row(titles):
            continue
        for asset_type, (title, source_url) in zip(types, cells):
            if not title:
                continue
            seen[asset_type].add(title)
            existing = (
                db.query(Asset)
                .filter(Asset.type == asset_type, Asset.title == title)
                .first()
            )
            if existing and existing.edited:
                continue
            file_status = "gap" if not source_url else "pending"
            payload = {
                "type": asset_type,
                "title": title,
                "url": source_url or None,
                "source_url": source_url,
                "module_ids": infer_module_ids(asset_type, title),
                "matrix_ref": infer_matrix_ref(asset_type, title),
                "file_status": file_status,
                "audiences": "",
                "status": "exists" if source_url else "gap",
                "notes": "",
            }
            if existing:
                stored = existing.file_status in {"stored", "processed", "preview"}
                for key, value in payload.items():
                    if stored and key in {"file_status", "url"}:
                        continue
                    setattr(existing, key, value)
                existing.edited = False
            else:
                db.add(Asset(**payload, edited=False))
            count += 1
    for asset_type, titles in seen.items():
        stale = (
            db.query(Asset)
            .filter(Asset.type == asset_type, Asset.edited.is_(False))
            .all()
        )
        for asset in stale:
            if asset.title not in titles and " — " not in asset.title:
                db.delete(asset)
    return count


def _sheet_rows(workbook, name: str) -> list[tuple]:
    sheet = workbook[name]
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(tuple(row))
    return rows


def _asset_sheet_rows(path: Path) -> list[list[tuple[str, str]]]:
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook["Videos  Photos  Reports"]
        rows: list[list[tuple[str, str]]] = []
        for row in sheet.iter_rows(min_row=2, max_col=5):
            drive_video, photo, report = row[1], row[3], row[4]

            video_title = "" if drive_video.value is None else str(drive_video.value).strip()
            video_link = (
                str(drive_video.hyperlink.target).strip()
                if drive_video.hyperlink and drive_video.hyperlink.target
                else ""
            )
            if "drive.google.com" not in video_link.lower() and "docs.google.com" not in video_link.lower():
                video_title = ""
                video_link = ""

            cells = [(video_title, video_link)]
            for cell in (photo, report):
                title = "" if cell.value is None else str(cell.value).strip()
                link = (
                    str(cell.hyperlink.target).strip()
                    if cell.hyperlink and cell.hyperlink.target
                    else ""
                )
                cells.append((title, link))
            rows.append(cells)
        return rows
    finally:
        workbook.close()


def seed_deck_vision(
    db: Session,
    rows: list[tuple],
    *,
    classify_with_llm: bool = False,
) -> int:
    """Seed ``DeckVisionRow`` from the Deck - Vision Mapping sheet."""
    from backend.pipeline.vision_deck import (
        classify_logline,
        classify_logline_with_llm,
        instructions_to_json,
        parse_instructions_json,
        parse_vision_sheet_rows,
    )

    parsed = parse_vision_sheet_rows(rows)
    if not parsed:
        return 0

    seen_keys: set[tuple[str, int]] = set()
    count = 0
    for row in parsed:
        key = (row.use_case, row.section_order)
        seen_keys.add(key)
        existing = (
            db.query(DeckVisionRow)
            .filter(
                DeckVisionRow.use_case == row.use_case,
                DeckVisionRow.section_order == row.section_order,
            )
            .first()
        )
        # LLM classifications are cached per logline; rule-based ones are cheap
        # and always recomputed so classifier fixes take effect on reseed.
        cached = (
            parse_instructions_json(existing.instructions_json or "")
            if existing is not None and (existing.logline or "") == row.logline
            else []
        )
        if cached and any(item.source == "llm" for item in cached):
            instructions_json = existing.instructions_json
        elif row.logline:
            classifier = classify_logline_with_llm if classify_with_llm else classify_logline
            instructions_json = instructions_to_json(classifier(row.logline))
        else:
            instructions_json = "[]"

        payload = {
            "section": row.section,
            "section_order": row.section_order,
            "use_case": row.use_case,
            "pages_json": json.dumps(row.pages, ensure_ascii=False),
            "section_pages_json": json.dumps(row.section_pages, ensure_ascii=False),
            "needs_more": row.needs_more,
            "logline": row.logline,
            "instructions_json": instructions_json,
            "design_status": row.design_status,
            "slide_status": row.slide_status,
        }
        if existing:
            for field_name, value in payload.items():
                setattr(existing, field_name, value)
        else:
            db.add(DeckVisionRow(**payload))
        count += 1

    for stale in db.query(DeckVisionRow).all():
        if (stale.use_case, stale.section_order) not in seen_keys:
            db.delete(stale)
    return count


def _vision_sheet_rows(workbook) -> list[tuple]:
    from backend.pipeline.vision_deck import VISION_SHEET_NAME

    if VISION_SHEET_NAME in workbook.sheetnames:
        return _sheet_rows(workbook, VISION_SHEET_NAME)
    for name in workbook.sheetnames:
        if name.strip() == VISION_SHEET_NAME.strip():
            return _sheet_rows(workbook, name)
    return []


def run_seed(
    xlsx_path: Path | None = None,
    db: Session | None = None,
    *,
    classify_vision_with_llm: bool = False,
) -> SeedCounts:
    path = Path(xlsx_path or DEFAULT_XLSX)
    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")
    close = False
    if db is None:
        Base.metadata.create_all(bind=engine)
        ensure_schema()
        db = SessionLocal()
        close = True
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        counts = SeedCounts(
            modules=seed_modules(db, _sheet_rows(workbook, "02 Modules")),
            facts=seed_facts(db, _sheet_rows(workbook, "08 Locked Facts")),
            recipes=seed_recipes(db, _sheet_rows(workbook, "03 Script Matrix")),
            objections=seed_objections(db, _sheet_rows(workbook, "09 Objection Pack")),
            assets=seed_assets(db, _asset_sheet_rows(path)),
            vision_rows=seed_deck_vision(
                db,
                _vision_sheet_rows(workbook),
                classify_with_llm=classify_vision_with_llm,
            ),
        )
        db.commit()
        return counts
    except Exception:
        db.rollback()
        raise
    finally:
        workbook.close()
        if close:
            db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Seed Pitch Studio from the One Story workbook")
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    parser.add_argument(
        "--classify-vision-llm",
        action="store_true",
        help="Use the LLM to classify vision-mapping loglines (default: deterministic rules)",
    )
    args = parser.parse_args()
    counts = run_seed(Path(args.xlsx), classify_vision_with_llm=args.classify_vision_llm)
    print(
        f"Seeded {counts.modules} modules, {counts.facts} facts, "
        f"{counts.recipes} recipes, {counts.objections} objections, "
        f"{counts.assets} assets, {counts.vision_rows} vision rows"
    )


if __name__ == "__main__":
    main()
